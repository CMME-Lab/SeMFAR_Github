"""Recalculate manuscript Table 3 with the boundary-refined ROI rule."""
import csv
import importlib.util
import json
import shutil
from pathlib import Path

import numpy as np
import torch
from docx import Document
from scipy import ndimage as ndi
from skimage.morphology import disk
from skimage.segmentation import morphological_chan_vese

ROOT = Path(r"D:\Woojeong_SR")
R50 = ROOT / "Comparison_300epochs_Data_ver10_50paired_1018unpaired_2"
R75 = ROOT / "Comparison_300epochs_Data_ver12_75paired_2"
R100 = ROOT / "Comparison_300epochs_Data_ver11_100paired_3"
PAIRED = ROOT / "data10" / "Paired_ver10"
OUT_ROOT = R100
MODELS = [
    (50, "Supervised-grad02", R50 / "Supervised_EDSR_Gradient_Lambda02_Adam_LR1e3"),
    (50, "EMA-re-epoch", R50 / "EMA_rrad02_re_epoch_AdamW_LR1e3"),
    (50, "Affine Residual", R50 / "Affine_Residual_Modulation_Adam"),
    (75, "Supervised-grad02", R75 / "Supervised_EDSR_Gradient_Lambda02_Adam_LR1e3"),
    (75, "EMA-re-epoch", R75 / "EMA_rrad02_re_epoch_AdamW_LR1e3"),
    (75, "Affine Residual", R75 / "Affine_Residual_Modulation_Adam"),
    (100, "Supervised-grad02", R100 / "Supervised_EDSR_Gradient_Lambda02_Adam_LR1e3"),
    (100, "EMA-re-epoch", R100 / "EMA_rrad02_re_epoch_AdamW_LR1e3"),
    (100, "Affine Residual", R100 / "Affine_Residual_Modulation_Adam"),
]

spec = importlib.util.spec_from_file_location("roi_eval", R50 / "recalculate_image_derived_roi_metrics.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)


def setup(gt):
    threshold = m.otsu(ndi.gaussian_filter(gt, 2.0))
    coarse = m.tumor_mask(gt, threshold)
    ys, xs = np.where(coarse); pad = 10
    box = (max(0, ys.min()-pad), min(gt.shape[0], ys.max()+pad+1),
           max(0, xs.min()-pad), min(gt.shape[1], xs.max()+pad+1))
    y0, y1, x0, x1 = box
    init = ndi.binary_erosion(coarse[y0:y1, x0:x1], structure=disk(3))
    if not init.any(): init = coarse[y0:y1, x0:x1]
    return box, init


def refine(image, box, init):
    y0, y1, x0, x1 = box
    crop = ndi.gaussian_filter(image[y0:y1, x0:x1].astype(np.float32), 1.0)
    level = morphological_chan_vese(crop, 160, init_level_set=init,
                                    smoothing=2, lambda1=1.0, lambda2=1.0)
    labels, count = ndi.label(level)
    if count:
        overlap = [np.logical_and(labels == idx, init).sum() for idx in range(1, count+1)]
        level = labels == (int(np.argmax(overlap))+1)
    level = ndi.binary_fill_holes(ndi.binary_closing(level, structure=disk(2)))
    result = np.zeros_like(image, dtype=bool); result[y0:y1, x0:x1] = level
    return result


def backup(path):
    if path.exists():
        target = path.with_name(path.stem + "_fixed_threshold_backup_20260930" + path.suffix)
        if not target.exists(): shutil.copy2(path, target)


def main():
    names = json.loads((R50 / "Affine_Residual_Modulation_Adam" / "split.json").read_text())["test"]
    cache = {}
    for name in names:
        gt = m.gray(PAIRED / f"{name}_HR.png"); box, init = setup(gt)
        cache[name] = (gt, box, init, refine(gt, box, init))
    per_image, summaries = [], []
    for index, (pairs, display, directory) in enumerate(MODELS, 1):
        rows = []
        for name in names:
            gt, box, init, gt_mask = cache[name]
            pred = m.gray(directory / "test_predictions" / f"{name}_pred.png")
            pred_mask = refine(pred, box, init)
            values = m.metrics(pred, gt, gt_mask, pred_mask)
            row = {"paired_cases": pairs, "model": display, "file": name,
                   "roi_psnr": values["roi_psnr"], "roi_ssim": values["roi_ssim"],
                   "roi_dice": values["dice"]}
            rows.append(row); per_image.append(row)
        checkpoint = torch.load(directory / "best.pt", map_location="cpu", weights_only=False)
        result = {"paired_cases": pairs, "model": display, "n": len(rows),
                  "best_epoch": int(checkpoint["epoch"])}
        for key in ("roi_psnr", "roi_ssim", "roi_dice"):
            values = np.asarray([r[key] for r in rows])
            result[key+"_mean"] = float(values.mean()); result[key+"_std"] = float(values.std(ddof=1))
        summaries.append(result); print(f"{index}/9 {pairs} {display}", flush=True)

    summary_csv = OUT_ROOT / "selected_models_50_75_100_comparison_300epoch_2.csv"
    detail_csv = OUT_ROOT / "selected_models_50_75_100_test10_roi_metrics_300epoch_2.csv"
    docx = OUT_ROOT / "selected_models_50_75_100_comparison_300epoch_2.docx"
    for path in (summary_csv, detail_csv, docx): backup(path)
    with detail_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(per_image[0])); w.writeheader(); w.writerows(per_image)
    output_rows = []
    for s in summaries:
        output_rows.append({"Paired cases": s["paired_cases"], "Model": s["model"],
            "ROI PSNR (dB)": f'{s["roi_psnr_mean"]:.4f} ± {s["roi_psnr_std"]:.4f}',
            "ROI SSIM": f'{s["roi_ssim_mean"]:.4f} ± {s["roi_ssim_std"]:.4f}',
            "ROI Dice": f'{s["roi_dice_mean"]:.4f} ± {s["roi_dice_std"]:.4f}',
            "Best epoch": s["best_epoch"]})
    with summary_csv.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(output_rows[0])); w.writeheader(); w.writerows(output_rows)
    (OUT_ROOT / "selected_models_50_75_100_comparison_300epoch_2_boundary_refined.json").write_text(
        json.dumps({"roi_method": "shared HR-localized window + per-image morphological Chan-Vese", "results": summaries}, indent=2), encoding="utf-8")
    doc = Document(); doc.add_heading("50 vs 75 vs 100 Paired Cases", 1)
    doc.add_paragraph("Mean ± sample standard deviation over the same ten test images. ROI masks use a shared HR-localized search window followed by per-image morphological Chan–Vese boundary refinement.")
    table = doc.add_table(rows=1, cols=6); table.style = "Table Grid"; headers = list(output_rows[0])
    for c,h in zip(table.rows[0].cells,headers): c.text=h
    for row in output_rows:
        for c,h in zip(table.add_row().cells,headers): c.text=str(row[h])
    doc.save(docx); doc.save(OUT_ROOT / "selected_models_50_75_100_comparison_300epoch_2_boundary_refined.docx")
    print(json.dumps(output_rows, indent=2, ensure_ascii=False), flush=True)


if __name__ == "__main__": main()
